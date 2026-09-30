import io
from openpyxl import load_workbook
from conftest import build_test_workbook
from app.extensions import db
from app.models import DSF, DSFValue, ImportColumn, User, AuditLog
from app.services.activity_service import activity_groups
from app.services.value_codec import serialize_value


def workbook():
    book=load_workbook(build_test_workbook())
    sheet=book.active
    sheet.insert_cols(10,2)
    sheet.cell(1,10,"Sous-Branche d'activité")
    sheet.cell(1,11,"Sous-Branche d'activité")
    for row in (2,3):
        sheet.cell(row,10,'A01001');sheet.cell(row,11,'Agriculture')
    result=io.BytesIO();book.save(result);book.close();result.seek(0)
    return result


def seed(client):
    response=client.post('/import/',data={'file':(workbook(),'activite.xlsx')})
    assert '/admin/' in response.location
    users=[User(username='controleur_un',role='controller'),User(username='controleur_deux',role='controller')]
    for user in users: user.set_password('test-password-activity')
    db.session.add_all(users);db.session.commit()
    return users,DSF.query.order_by(DSF.id).all()


def test_duplicate_activity_headers_existing_data_and_live_corrections(client):
    users,rows=seed(client)
    columns=ImportColumn.query.filter_by(variable_name="Sous-Branche d'activité").order_by(ImportColumn.column_index).all()
    assert [c.column_index for c in columns]==[10,11]
    activity,groups=activity_groups()
    assert len(groups)==1 and groups[0]['free']==2
    assert activity[rows[0].id]['label']=='A01001 - Agriculture'
    # Grouping reflects current DB values, without reimporting or a stale cache.
    cell=DSFValue.query.filter_by(dsf_id=rows[0].id,import_column_id=columns[1].id).one()
    original=cell.original_value
    cell.current_value=serialize_value('Culture');db.session.commit()
    assert len(activity_groups()[1])==2 and cell.original_value==original
    html=client.get('/admin/').get_data(as_text=True)
    assert 'A01001 - Culture' in html and 'A01001 - Agriculture' in html


def test_branch_assignment_scope_and_repeat(client):
    users,rows=seed(client)
    sid=rows[0].import_session_id
    rows[0].assignee=users[1];rows[0].status='completed';db.session.commit()
    client.post('/import/',data={'file':(workbook(),'autre.xlsx')})
    branch=activity_groups(sid)[1][0]['key']
    url=f'/admin/import-sessions/{sid}/assign-branch'
    response=client.post(url,data={'user_id':users[0].id,'branch':branch})
    assert response.status_code==302
    db.session.expire_all()
    assert rows[0].assigned_to_id==users[1].id and rows[0].status=='completed'
    assert rows[1].assigned_to_id==users[0].id
    assert DSF.query.filter(DSF.import_session_id!=sid,DSF.assigned_to_id.is_not(None)).count()==0
    assert AuditLog.query.filter_by(action='affectation par sous-branche').count()==1
    client.post(url,data={'user_id':users[1].id,'branch':branch})
    assert AuditLog.query.filter_by(action='affectation par sous-branche').count()==1
    assert client.post(url,data={'user_id':users[0].id,'branch':'invalid'}).status_code==400


def test_controller_global_search_is_read_only_and_assigned_only(client):
    users,rows=seed(client)
    rows[0].assignee=users[1];db.session.commit()
    with client.session_transaction() as session: session['user_id']=users[0].id
    for args in ({},{'assignment':'unassigned'},{'q':'Agriculture'}):
        response=client.get('/admin/search',query_string=args)
        assert response.status_code==200
        html=response.get_data(as_text=True)
        assert f'data-dsf-id="{rows[0].id}"' in html
        assert f'data-dsf-id="{rows[1].id}"' not in html
        assert '>Ouvrir</a>' not in html
        assert 'Recherche globale' in html
    assert client.get(f'/dsf/{rows[0].id}').status_code == 403
    assert client.post(f'/admin/import-sessions/{rows[0].import_session_id}/assign-branch',data={'user_id':users[0].id,'branch':'[]'}).status_code==403
    assert client.get('/admin/').status_code==403


def test_missing_branch_and_exact_filter(client,imported_session):
    groups=activity_groups(imported_session.id)[1]
    assert groups[0]['key']=='[]' and groups[0]['free']==2
    seed(client)
    response=client.get('/admin/search',query_string={'branch':'[]'})
    assert response.get_data(as_text=True).count('data-dsf-id=')==2
    assert 'Aucune DSF' in client.get('/admin/search',query_string={'branch':'invalid'}).get_data(as_text=True)
