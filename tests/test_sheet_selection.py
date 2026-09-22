import io
from openpyxl import load_workbook
from conftest import build_test_workbook
from app.extensions import db
from app.models import ImportSession, DSF, ImportColumn, User


def multi_sheet():
    workbook=load_workbook(build_test_workbook())
    first=workbook.active;first.title='Première'
    second=workbook.copy_worksheet(first);second.title='DSF à contrôler'
    second.cell(2,1,'CHOIX-002');second.insert_rows(1,2)
    output=io.BytesIO();workbook.save(output);workbook.close();output.seek(0)
    return output


def test_lists_sheets_without_importing(client,app):
    response=client.post('/import/sheets',data={'file':(multi_sheet(),'multi.xlsx')})
    assert response.status_code==200
    assert response.json['sheets']==['Première','DSF à contrôler']
    assert ImportSession.query.count()==0
    assert not list(app.config['UPLOAD_FOLDER'].rglob('*.xlsx'))


def test_imports_only_chosen_sheet_with_original_headers(client):
    response=client.post('/import/',data={'file':(multi_sheet(),'multi.xlsx'),'sheet_name':'DSF à contrôler'})
    assert response.status_code==302
    session=ImportSession.query.one()
    assert session.sheet_name=='DSF à contrôler'
    assert session.header_row==3
    assert DSF.query.count()==2
    assert DSF.query.order_by(DSF.row_index).first().numero_dsf=='CHOIX-002'
    assert ImportColumn.query.filter_by(variable_name='Ville').count()==2


def test_invalid_selection_and_corrupt_file_do_not_import(client):
    response=client.post('/import/sheets',data={'file':(io.BytesIO(b'bad'),'bad.xlsx')})
    assert response.status_code==400
    response=client.post('/import/',data={'file':(multi_sheet(),'multi.xlsx'),'sheet_name':'absente'},follow_redirects=True)
    assert response.status_code==200
    assert ImportSession.query.count()==0


def test_preview_admin_and_csrf_protected(app,client):
    app.config['TESTING']=False
    assert client.post('/import/sheets',data={'file':(multi_sheet(),'multi.xlsx')}).status_code==400
    user=User(username='testcontroleur',role='controller');user.set_password('motdepasse-test')
    db.session.add(user);db.session.commit()
    with client.session_transaction() as s:s['user_id']=user.id;token=s['csrf_token']
    assert client.post('/import/sheets',data={'file':(multi_sheet(),'multi.xlsx'),'csrf_token':token}).status_code==403
