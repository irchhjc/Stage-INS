import io

from openpyxl import Workbook, load_workbook

from app.models import User
from app.services.auth_service import create_user


def _workbook(rows, headers=("NOM COMPLET", "NOM UTILISATEUR", "MOT DE PASSE")):
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(headers)
    for row in rows:
        worksheet.append(row)
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    output.seek(0)
    return output


def test_admin_imports_valid_controllers_and_reports_rejected_rows(client):
    create_user("existant", "motdepasse-existant", full_name="Compte existant")
    source = _workbook(
        [
            ("Alice Martin", "alice.martin", "motdepasse-alice"),
            ("Boris Nsimba", "boris_nsimba", "motdepasse-boris"),
            ("Compte existant", "existant", "motdepasse-autre"),
            ("Nom invalide", "NomMajuscule", "motdepasse-valide"),
            ("Mot court", "mot.court", "court"),
            (None, None, None),
        ]
    )

    response = client.post(
        "/admin/users/import",
        data={"file": (source, "controleurs.xlsx")},
        content_type="multipart/form-data",
        follow_redirects=True,
    )

    assert response.status_code == 200
    html = response.get_data(as_text=True)
    assert "2 compte(s) contrôleur créé(s)" in html
    assert "3 ligne(s) rejetée(s)" in html
    assert "1 ligne(s) vide(s) ignorée(s)" in html
    alice = User.query.filter_by(username="alice.martin").one()
    assert alice.full_name == "Alice Martin"
    assert alice.role == "controller"
    assert alice.password_hash != "motdepasse-alice"
    assert alice.check_password("motdepasse-alice")


def test_controller_import_requires_all_headers_and_creates_nothing(client):
    source = _workbook(
        [("Alice Martin", "alice.martin")],
        headers=("NOM COMPLET", "NOM UTILISATEUR"),
    )

    response = client.post(
        "/admin/users/import",
        data={"file": (source, "controleurs.xlsx")},
        content_type="multipart/form-data",
        follow_redirects=True,
    )

    assert "MOT DE PASSE" in response.get_data(as_text=True)
    assert User.query.filter_by(username="alice.martin").first() is None


def test_controller_import_rejects_non_excel_file(client):
    response = client.post(
        "/admin/users/import",
        data={"file": (io.BytesIO(b"not an excel file"), "controleurs.csv")},
        content_type="multipart/form-data",
        follow_redirects=True,
    )

    assert "Seuls les fichiers Excel .xlsx sont acceptés" in response.get_data(as_text=True)


def test_admin_can_download_controller_import_template(client):
    response = client.get("/admin/users/import-template")

    assert response.status_code == 200
    assert response.headers["Content-Disposition"].startswith("attachment;")
    workbook = load_workbook(io.BytesIO(response.data), read_only=True)
    try:
        assert tuple(cell.value for cell in next(workbook.active.iter_rows())) == (
            "NOM COMPLET",
            "NOM UTILISATEUR",
            "MOT DE PASSE",
        )
    finally:
        workbook.close()
