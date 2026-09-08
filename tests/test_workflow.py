import io

from openpyxl import load_workbook
from sqlalchemy import text

from app.extensions import db
from app.models import AuditLog, DSF, DSFValue, FicheStatus, ImportColumn, User
from app.services.dsf_service import FINAL_FICHE_STATUSES, mark_fiche_not_provided


def complete_dsf(dsf, operator="Testeur"):
    unfinished = [fiche for fiche in dsf.fiche_statuses if fiche.status not in FINAL_FICHE_STATUSES]
    for fiche in unfinished:
        mark_fiche_not_provided(dsf, fiche.fiche_code, operator)
    db.session.refresh(dsf)
    assert dsf.status == "completed"


def test_institutional_branding_and_logo_assets(client):
    dashboard = client.get("/")
    html = dashboard.get_data(as_text=True)
    assert dashboard.status_code == 200
    assert "Institut National de la Statistique" in html
    assert "Banque mondiale" in html
    assert "HISWACA" in html
    assert client.get("/brand-assets/ins_congo.png").status_code == 200
    assert client.get("/brand-assets/partenaire1.jpg").mimetype == "image/jpeg"


def test_sqlite_waits_for_transient_writes(app):
    with app.app_context():
        busy_timeout = db.session.execute(text("PRAGMA busy_timeout")).scalar()
    assert busy_timeout == 60000


def test_login_creates_a_permanent_session_and_api_expiry_is_explicit(app):
    anonymous_client = app.test_client()
    login = anonymous_client.post(
        "/auth/login",
        data={"username": "irch", "password": "15081960irchdefluviaire"},
    )
    assert login.status_code == 302
    with anonymous_client.session_transaction() as user_session:
        assert user_session.permanent is True

    anonymous_client.post("/auth/logout")
    expired_api = anonymous_client.patch("/dsf/api/values/1", json={"value": "1"})
    assert expired_api.status_code == 401
    assert expired_api.get_json() == {"ok": False, "error": "Authentification requise."}


def test_import_preserves_duplicate_headers_and_duplicate_niu(client, imported_session):
    assert imported_session.row_count == 2
    assert DSF.query.filter_by(niu="M001").count() == 2
    cities = ImportColumn.query.filter_by(import_session_id=imported_session.id, variable_name="Ville").all()
    assert len(cities) == 2
    assert cities[0].column_index != cities[1].column_index
    response = client.get(f"/?session_id={imported_session.id}")
    assert response.status_code == 200
    assert b"ALPHA SARL" in response.data
    dsf = DSF.query.first()
    fiche_page = client.get(f"/dsf/{dsf.id}/fiche/BILAN_ACTIF")
    fiche_html = fiche_page.get_data(as_text=True)
    assert "previousFicheLink" in fiche_html
    assert "nextFicheLink" in fiche_html
    assert "Suivante : Bilan Passif" in fiche_html
    assert "toggleFicheSidebar" in fiche_html
    assert "toggleAppHeader" in fiche_html
    assert "toggleFocusMode" in fiche_html


def test_edit_validate_search_and_export(client, imported_session):
    dsf = DSF.query.order_by(DSF.id).first()
    target = (
        DSFValue.query.join(ImportColumn)
        .filter(DSFValue.dsf_id == dsf.id, ImportColumn.variable_name == "IMMOBILISATIONS INCORPORELLES (NET_N)")
        .one()
    )
    edit = client.patch(
        f"/dsf/api/values/{target.id}",
        json={"value": "81", "status": "verified", "operator": "Testeur"},
    )
    assert edit.status_code == 200
    assert edit.get_json()["corrected"] is True
    assert AuditLog.query.filter_by(dsf_id=dsf.id, action="correction").count() == 1

    blocked_validation = client.post(
        f"/dsf/api/{dsf.id}/fiches/BILAN_ACTIF/validate",
        json={"operator": "Testeur"},
    )
    assert blocked_validation.status_code == 409
    assert blocked_validation.get_json()["requires_confirmation"] is True
    assert blocked_validation.get_json()["anomaly_count"] >= 1

    confirmed_validation = client.post(
        f"/dsf/api/{dsf.id}/fiches/BILAN_ACTIF/validate",
        json={"operator": "Testeur", "acknowledge_anomalies": True},
    )
    assert confirmed_validation.status_code == 200
    assert FicheStatus.query.filter_by(dsf_id=dsf.id, fiche_code="BILAN_ACTIF").one().status == "verified"
    assert AuditLog.query.filter_by(
        dsf_id=dsf.id,
        fiche_code="BILAN_ACTIF",
        action="validation fiche avec anomalies",
    ).count() == 1

    search = client.get(f"/dsf/api/{dsf.id}/search?q=Clients")
    assert search.status_code == 200
    assert search.get_json()["ok"] is True

    validation = client.post(
        f"/dsf/api/{dsf.id}/fiches/IDENT/validate",
        json={"operator": "Testeur"},
    )
    assert validation.status_code == 200
    assert FicheStatus.query.filter_by(dsf_id=dsf.id, fiche_code="IDENT").one().status == "verified"

    niu_value = (
        DSFValue.query.join(ImportColumn)
        .filter(DSFValue.dsf_id == dsf.id, ImportColumn.variable_name == "NIU")
        .one()
    )
    reopened_by_edit = client.patch(
        f"/dsf/api/values/{niu_value.id}",
        json={"value": "M001", "status": "verified", "operator": "Testeur"},
    )
    assert reopened_by_edit.status_code == 200
    assert FicheStatus.query.filter_by(dsf_id=dsf.id, fiche_code="IDENT").one().status == "in_progress"

    complete_dsf(dsf)
    exported = client.post(f"/export/{imported_session.id}")
    assert exported.status_code == 200
    workbook = load_workbook(io.BytesIO(exported.data), data_only=False)
    worksheet = workbook[imported_session.sheet_name]
    assert worksheet.cell(dsf.row_index, target.column.column_index).value == 81
    assert worksheet.cell(dsf.row_index, target.column.column_index).font.color.rgb == "FF008000"
    assert worksheet.cell(dsf.row_index, target.column.column_index).fill.fgColor.rgb == "FFE2F0D9"
    assert "JOURNAL_CONTROLE" in workbook.sheetnames
    assert worksheet.cell(1, target.column.column_index).value == target.variable_name
    workbook.close()


def test_export_contains_only_fully_controlled_dsfs(client, imported_session):
    dashboard = client.get(f"/?session_id={imported_session.id}")
    assert "Exporter toutes les DSF terminées (0)" in dashboard.get_data(as_text=True)

    no_completed = client.post(f"/export/{imported_session.id}", follow_redirects=True)
    assert no_completed.status_code == 200
    assert "Aucune DSF entièrement contrôlée" in no_completed.get_data(as_text=True)

    first, second = DSF.query.order_by(DSF.id).all()
    complete_dsf(first)
    assert second.status == "not_started"

    exported = client.post(f"/export/{imported_session.id}")
    assert exported.status_code == 200
    workbook = load_workbook(io.BytesIO(exported.data), data_only=False)
    worksheet = workbook[imported_session.sheet_name]
    numero_column = ImportColumn.query.filter_by(
        import_session_id=imported_session.id,
        variable_name="NUMERO DE LA DSF",
    ).one()
    exported_numbers = [
        worksheet.cell(row=row_index, column=numero_column.column_index).value
        for row_index in range(imported_session.header_row + 1, worksheet.max_row + 1)
    ]
    assert exported_numbers == [first.numero_dsf]
    assert second.numero_dsf not in exported_numbers
    workbook.close()


def test_invalid_workbook_is_rejected_without_crash(client, app):
    response = client.post(
        "/import/",
        data={"file": (io.BytesIO(b"not-an-xlsx"), "bad.xlsx")},
        content_type="multipart/form-data",
        follow_redirects=True,
    )
    assert response.status_code == 200
    assert "classeur .xlsx lisible" in response.get_data(as_text=True)
    assert not list(app.config["UPLOAD_FOLDER"].rglob("*.xlsx"))


def test_admin_creates_assigns_and_restricts_controller_access(client, imported_session):
    invalid_username = client.post(
        "/admin/users",
        data={"username": "Controleur", "password": "motdepassevalide"},
        follow_redirects=True,
    )
    assert "entièrement en minuscules" in invalid_username.get_data(as_text=True)
    assert User.query.filter_by(username="Controleur").first() is None

    short_password = client.post(
        "/admin/users",
        data={"username": "controleur", "password": "tropcourt"},
        follow_redirects=True,
    )
    assert "au moins 10 caractères" in short_password.get_data(as_text=True)

    created = client.post(
        "/admin/users",
        data={"username": "controleur", "password": "motdepasse10"},
        follow_redirects=True,
    )
    assert created.status_code == 200
    controller = User.query.filter_by(username="controleur").one()
    assert controller.password_hash != "motdepasse10"

    first, second = DSF.query.order_by(DSF.id).all()
    assignment = client.post(
        f"/admin/dsfs/{first.id}/assign",
        data={"user_id": controller.id},
        follow_redirects=True,
    )
    assert assignment.status_code == 200
    assert db.session.get(DSF, first.id).assigned_to_id == controller.id
    assert "Affectation verrouillée" in assignment.get_data(as_text=True)
    reassignment = client.post(
        f"/admin/dsfs/{first.id}/assign",
        data={"user_id": ""},
        follow_redirects=True,
    )
    assert "déjà affectée" in reassignment.get_data(as_text=True)
    assert db.session.get(DSF, first.id).assigned_to_id == controller.id
    assert AuditLog.query.filter_by(dsf_id=first.id, action="affectation DSF").count() == 1

    client.post(
        "/admin/users",
        data={"username": "controleur2", "password": "motdepasse20"},
    )
    controller2 = User.query.filter_by(username="controleur2").one()
    bulk_assignment = client.post(
        f"/admin/import-sessions/{imported_session.id}/assign-all",
        data={"user_id": controller2.id},
        follow_redirects=True,
    )
    bulk_html = bulk_assignment.get_data(as_text=True)
    assert "1 DSF affectée(s) à controleur2" in bulk_html
    assert "1 DSF déjà affectée(s)" in bulk_html
    assert db.session.get(DSF, first.id).assigned_to_id == controller.id
    assert db.session.get(DSF, second.id).assigned_to_id == controller2.id
    assert AuditLog.query.filter_by(dsf_id=second.id, action="affectation DSF groupée").count() == 1

    client.post("/auth/logout")
    anonymous = client.get("/", follow_redirects=False)
    assert anonymous.status_code == 302
    assert "/auth/login" in anonymous.headers["Location"]
    expired_assignment = client.post(
        f"/admin/dsfs/{second.id}/assign",
        data={"user_id": controller.id},
        follow_redirects=False,
    )
    assert expired_assignment.status_code == 302
    assert "/auth/login" in expired_assignment.headers["Location"]
    logged_in = client.post(
        "/auth/login",
        data={"username": "controleur", "password": "motdepasse10"},
        follow_redirects=False,
    )
    assert logged_in.status_code == 302

    dashboard = client.get("/")
    html = dashboard.get_data(as_text=True)
    assert "DSF-001" in html
    assert "DSF-002" not in html
    assert client.get(f"/dsf/{first.id}").status_code == 200
    assert client.get(f"/dsf/{second.id}").status_code == 403
    assert client.get("/admin/").status_code == 403
    assert client.get("/import/").status_code == 403

    target = DSFValue.query.filter_by(dsf_id=first.id).first()
    edited = client.patch(
        f"/dsf/api/values/{target.id}",
        json={"value": "DSF-001-C", "status": "verified", "operator": "faux-operateur"},
    )
    assert edited.status_code == 200
    latest_log = AuditLog.query.filter_by(dsf_id=first.id).order_by(AuditLog.id.desc()).first()
    assert latest_log.operator == "controleur"

    denied_target = DSFValue.query.filter_by(dsf_id=second.id).first()
    assert client.patch(
        f"/dsf/api/values/{denied_target.id}",
        json={"value": "interdit", "status": "verified"},
    ).status_code == 403

    complete_dsf(first, operator="controleur")
    exported = client.post(f"/export/{imported_session.id}")
    assert exported.status_code == 200
    workbook = load_workbook(io.BytesIO(exported.data), data_only=False)
    worksheet = workbook[imported_session.sheet_name]
    assert worksheet.max_row == imported_session.header_row + 1
    assert "JOURNAL_CONTROLE" in workbook.sheetnames
    workbook.close()
