from copy import copy
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from app.extensions import db
from app.models import DSF, DSFValue, ImportColumn
from app.services.export_service import _retain_rows, export_controlled_workbook
from app.services.value_codec import serialize_value


def test_sparse_row_compaction_matches_previous_behavior():
    book=Workbook();sheet=book.active
    for row in (1,2,4,6,8):
        sheet.cell(row,1,f'row-{row}')
        sheet.cell(row,3,f'=A{row}')
    sheet.cell(6,3).font=Font(name='Courier New',italic=True)
    old=book.copy_worksheet(sheet)
    for row in range(old.max_row,2,-1):
        if row not in {4,6}:old.delete_rows(row)
    _retain_rows(sheet,2,{4,6})
    assert sheet.max_row==old.max_row
    for row in old:
        for cell in row:
            actual=sheet.cell(cell.row,cell.column)
            assert actual.value==cell.value
            assert copy(actual.font)==copy(cell.font)


def test_shared_styles_preserve_source_fonts_without_correction_leak(client,imported_session):
    first=DSF.query.order_by(DSF.row_index).first();first.status='completed'
    book=load_workbook(imported_session.filepath)
    sheet=book[imported_session.sheet_name]
    sheet.cell(first.row_index,3).font=Font(name='Courier New',size=17,italic=True)
    book.save(imported_session.filepath);book.close()
    values=DSFValue.query.join(ImportColumn).filter(DSFValue.dsf_id==first.id,ImportColumn.column_index.in_([3,4,5])).order_by(ImportColumn.column_index).all()
    for value in values:value.status='verified';value.verified=True
    values[1].corrected=True;values[1].current_value=serialize_value('CHANGED')
    db.session.commit()
    path=export_controlled_workbook(imported_session.id)
    book=load_workbook(path);sheet=book[imported_session.sheet_name]
    assert sheet.cell(2,3).font.name=='Courier New'
    assert sheet.cell(2,3).font.sz==17 and sheet.cell(2,3).font.i
    assert sheet.cell(2,3).font.color.rgb=='FF008000'
    assert sheet.cell(2,4).value=='CHANGED'
    assert sheet.cell(2,4).fill.fgColor.rgb=='FFE2F0D9'
    assert sheet.cell(2,5).fill.fgColor.rgb!='FFE2F0D9'
    assert book['JOURNAL_CONTROLE'].max_row==4
    book.close()


def test_not_started_export_streams_only_untouched_dsfs_from_the_database(client, imported_session):
    from app.services.export_service import export_not_started_workbook

    first, second = DSF.query.order_by(DSF.row_index).all()
    first.status = 'completed'
    db.session.commit()
    headers = [c.variable_name for c in ImportColumn.query.filter_by(import_session_id=imported_session.id).order_by(ImportColumn.column_index)]

    path = export_not_started_workbook(imported_session.id)
    book = load_workbook(path)
    sheet = book['DSF_NON_COMMENCEES']

    assert 'dsf_non_commencees' in path.name
    assert book.sheetnames == ['DSF_NON_COMMENCEES']
    assert [cell.value for cell in sheet[1]] == headers
    assert sheet.max_row == 2
    assert sheet.cell(2, 1).value == second.numero_dsf
    book.close()


def test_not_started_export_route_is_admin_only_and_reports_empty(client, imported_session):
    response = client.post(f'/export/{imported_session.id}/not-started')
    assert response.status_code == 200
    assert response.headers['Content-Disposition'].startswith('attachment')

    DSF.query.update({'status': 'in_progress'})
    db.session.commit()
    response = client.post(f'/export/{imported_session.id}/not-started')
    assert response.status_code == 302


def test_admin_page_offers_the_not_started_export_with_count(client, imported_session):
    html = client.get('/admin/').get_data(as_text=True)

    assert 'Exporter les DSF non commencées (2)' in html
    assert f'/export/{imported_session.id}/not-started' in html


def test_controller_cannot_export_not_started_dsfs(app, client, imported_session):
    from app.services.auth_service import create_user

    create_user('ctrl', 'motdepasse10', role='controller')
    other = app.test_client()
    other.post('/auth/login', data={'username': 'ctrl', 'password': 'motdepasse10'})

    assert other.post(f'/export/{imported_session.id}/not-started').status_code == 403
