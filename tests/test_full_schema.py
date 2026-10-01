import io
import re
from collections import Counter, defaultdict
from pathlib import Path
from types import SimpleNamespace

from openpyxl import Workbook, load_workbook

from app.config.fiche_mapping import FICHE_DEFINITIONS
from app.extensions import db
from app.models import DSF, DSFValue, ImportColumn, ImportSession
from app.services.mapping_service import build_accounting_sections, build_column_mapping, normalize_label
from app.services.schema_service import (
    FULL_DSF_COLUMN_COUNT,
    FULL_DSF_SCHEMA_CODE,
    identify_schema,
)
from app.services.value_codec import raw_input_value


HEADERS_PATH = Path(__file__).parent / "fixtures" / "new_dsf_headers_20260925.txt"


def _reference_headers():
    return HEADERS_PATH.read_text(encoding="utf-8-sig").strip().split("\t")


def _reference_values_by_fiche():
    values_by_fiche = defaultdict(list)
    for item in build_column_mapping(_reference_headers()):
        values_by_fiche[item["fiche_code"]].append(
            SimpleNamespace(variable_name=item["variable_name"])
        )
    return values_by_fiche


def test_bilan_actif_uses_one_complete_four_measure_table():
    sections = build_accounting_sections(_reference_values_by_fiche()["BILAN_ACTIF"])

    assert len(sections) == 1
    assert sections[0]["column_count"] == 4
    assert [slot["label"] for slot in sections[0]["slots"]] == [
        "BRUT",
        "AMORT./DÉPRÉC.",
        "NET N",
        "NET N-1",
    ]
    assert len(sections[0]["rows"]) == 29
    assert all(len(row["cells"]) == 4 for row in sections[0]["rows"])


def test_every_full_schema_table_contains_only_real_cells():
    for values in _reference_values_by_fiche().values():
        for section in build_accounting_sections(values):
            expected_slots = {slot["key"] for slot in section["slots"]}
            assert expected_slots
            assert all(set(row["cells"]) == expected_slots for row in section["rows"])


def test_question_variables_use_compact_blocks():
    headers = _reference_headers()
    expected_questions = [header.strip() for header in headers if header.strip().endswith("?")]
    rendered_questions = []

    for values in _reference_values_by_fiche().values():
        for section in build_accounting_sections(values):
            if section["is_question_block"]:
                assert section["column_count"] == 1
                assert [slot["label"] for slot in section["slots"]] == ["Valeur"]
                rendered_questions.extend(row["poste"] for row in section["rows"])

    assert len(expected_questions) == 31
    assert rendered_questions == expected_questions


def test_short_tables_use_natural_height_without_changing_their_columns():
    compact_sections = []

    for values in _reference_values_by_fiche().values():
        for section in build_accounting_sections(values):
            expected_compact = not section["is_question_block"] and len(section["rows"]) <= 3
            assert section["is_compact_table"] is expected_compact
            if expected_compact:
                compact_sections.append(section)
                expected_slots = {slot["key"] for slot in section["slots"]}
                assert all(set(row["cells"]) == expected_slots for row in section["rows"])

    assert compact_sections
    assert any(
        section["column_count"] == 1 and len(section["rows"]) == 2
        for section in compact_sections
    )


def test_note_27b_uses_one_complete_table_without_subtables():
    source_values = _reference_values_by_fiche()["NOTE_27B"]
    sections = build_accounting_sections(source_values, fiche_code="NOTE_27B")

    assert [section["title"] for section in sections] == [
        "Question de contrôle",
        "Tableau complet - effectifs et masse salariale",
    ]
    complete_table = sections[1]
    assert [slot["label"] for slot in complete_table["slots"]] == [
        "Hommes",
        "Femmes",
        "Total",
    ]
    assert len(complete_table["rows"]) == 22
    assert [row["poste"] for row in complete_table["rows"]][-3:] == [
        "Total - masse salariale (groupe 2)",
        "Total - effectifs (1+2)",
        "Total - masse salariale (1+2)",
    ]
    assert all(len(row["cells"]) == 3 for row in complete_table["rows"])

    rendered_variables = [
        value.variable_name
        for section in sections
        for value in _section_values_for_test(section)
    ]
    assert rendered_variables == [value.variable_name for value in source_values]


def test_measure_families_receive_explicit_business_titles():
    values_by_fiche = _reference_values_by_fiche()

    assert build_accounting_sections(values_by_fiche["BILAN_ACTIF"])[0]["title"] == (
        "Valeurs brutes, amortissements et valeurs nettes"
    )
    assert build_accounting_sections(values_by_fiche["BILAN_PASSIF"])[0]["title"] == (
        "Comparaison des exercices N et N-1"
    )

    note_3a_titles = {
        section["title"] for section in build_accounting_sections(values_by_fiche["NOTE_3A"])
    }
    assert "Mouvements des immobilisations brutes" in note_3a_titles
    assert "Mouvements des amortissements et dépréciations" in note_3a_titles

    amort_titles = {
        section["title"]
        for section in build_accounting_sections(values_by_fiche["NOTE_AMORT"])
    }
    assert "Cessions d'immobilisations" in amort_titles

    note_4_titles = {
        section["title"] for section in build_accounting_sections(values_by_fiche["NOTE_4"])
    }
    assert "Valeurs des exercices et échéancier par maturité" in note_4_titles


def test_complex_fiches_are_grouped_into_complete_compatible_tables():
    values_by_fiche = _reference_values_by_fiche()
    expected_tables = {
        "NOTE_3A": [(20, 7), (14, 4)],
        "NOTE_AMORT": [(15, 5)],
        "NOTE_4": [(12, 5)],
        "NOTE_16A": [(17, 5), (14, 2)],
        "NOTE_24": [(15, 2)],
        "NOTE_34": [(49, 2)],
    }

    for fiche_code, expected_shapes in expected_tables.items():
        sections = build_accounting_sections(
            values_by_fiche[fiche_code], fiche_code=fiche_code
        )
        tables = [section for section in sections if not section["is_question_block"]]
        assert [
            (len(section["rows"]), section["column_count"])
            for section in tables
        ] == expected_shapes


def _section_values_for_test(section):
    for row in section["rows"]:
        for slot in section["slots"]:
            value = row["cells"].get(slot["key"])
            if value is not None:
                yield value


def _reference_workbook(headers):
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "DSF complète"
    worksheet.append(headers)

    values = [None] * len(headers)
    identity_values = {
        "numero de la dsf": "DSF-COMPLETE-001",
        "numero_t": "T-COMPLETE-001",
        "niu": "NIU-COMPLETE-001",
        "raison sociale": "ENTREPRISE SCHEMA COMPLET",
        "sigle usuel": "ESC",
        "annee": 2026,
    }
    for index, header in enumerate(headers):
        normalized = normalize_label(header)
        if normalized in identity_values:
            values[index] = identity_values[normalized]
    worksheet.append(values)

    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    output.seek(0)
    return output


def test_full_1777_column_schema_is_recognized_imported_and_exported(client):
    headers = _reference_headers()
    profile = identify_schema(headers)

    assert len(headers) == FULL_DSF_COLUMN_COUNT
    assert profile["recognized"] is True
    assert profile["code"] == FULL_DSF_SCHEMA_CODE

    response = client.post(
        "/import/",
        data={"file": (_reference_workbook(headers), "dsf_complete_1777.xlsx")},
        follow_redirects=True,
    )
    html = response.get_data(as_text=True)
    assert response.status_code == 200
    assert "Schéma DSF complet - 1 777 colonnes reconnu" in html

    import_session = ImportSession.query.one()
    stored_columns = ImportColumn.query.order_by(ImportColumn.column_index).all()
    dsf = DSF.query.one()

    assert import_session.column_count == FULL_DSF_COLUMN_COUNT
    assert [column.variable_name for column in stored_columns] == headers
    assert [column.column_index for column in stored_columns] == list(
        range(1, FULL_DSF_COLUMN_COUNT + 1)
    )
    assert {column.fiche_code for column in stored_columns} == {
        definition["code"] for definition in FICHE_DEFINITIONS
    }
    assert DSFValue.query.filter_by(dsf_id=dsf.id).count() == FULL_DSF_COLUMN_COUNT
    assert dsf.numero_dsf == "DSF-COMPLETE-001"
    assert dsf.numero_dsf_t == "T-COMPLETE-001"
    assert dsf.niu == "NIU-COMPLETE-001"
    assert dsf.raison_sociale == "ENTREPRISE SCHEMA COMPLET"
    assert dsf.sigle == "ESC"
    assert dsf.annee == "2026"

    duplicate_header = next(name for name, count in Counter(headers).items() if count > 1)
    expected_positions = [index for index, name in enumerate(headers, start=1) if name == duplicate_header]
    stored_positions = [
        column.column_index for column in stored_columns if column.variable_name == duplicate_header
    ]
    assert stored_positions == expected_positions

    dsf.status = "completed"
    db.session.commit()
    exported = client.post(f"/export/{import_session.id}")
    assert exported.status_code == 200
    exported_workbook = load_workbook(io.BytesIO(exported.data), read_only=True, data_only=False)
    try:
        exported_headers = [
            cell.value
            for cell in next(
                exported_workbook[import_session.sheet_name].iter_rows(
                    min_row=import_session.header_row,
                    max_row=import_session.header_row,
                    max_col=FULL_DSF_COLUMN_COUNT,
                )
            )
        ]
        assert exported_headers == headers
    finally:
        exported_workbook.close()


def test_every_full_schema_variable_is_rendered_as_an_editable_input(client):
    headers = _reference_headers()
    response = client.post(
        "/import/",
        data={"file": (_reference_workbook(headers), "dsf_complete_editable.xlsx")},
        follow_redirects=False,
    )
    assert response.status_code == 302

    dsf = DSF.query.one()
    columns_by_fiche = defaultdict(list)
    for column in ImportColumn.query.order_by(ImportColumn.column_index):
        columns_by_fiche[column.fiche_code].append(column)

    rendered_value_ids = set()
    compact_question_rows = 0
    values_by_fiche = _reference_values_by_fiche()
    for definition in FICHE_DEFINITIONS:
        fiche_code = definition["code"]
        page = client.get(f"/dsf/{dsf.id}/fiche/{fiche_code}")
        assert page.status_code == 200
        html = page.get_data(as_text=True)
        inputs = re.findall(r'<input class="value-input"[^>]*>', html)
        assert len(inputs) == len(columns_by_fiche[fiche_code])
        assert all(" disabled" not in input_tag for input_tag in inputs)
        assert 'class="empty-value"' not in html
        compact_question_rows += html.count('class="compact-question-row"')
        expected_sections = build_accounting_sections(
            values_by_fiche[fiche_code],
            fiche_code=fiche_code,
        )
        expected_tables = sum(not section["is_question_block"] for section in expected_sections)
        assert html.count('class="table accounting-table') == expected_tables
        if fiche_code == "NOTE_27B":
            assert "Tableau complet - effectifs et masse salariale" in html
            assert "Tableau 2 -" not in html
        rendered_value_ids.update(
            int(value_id) for value_id in re.findall(r'data-value-id="(\d+)"', html)
        )

    stored_value_ids = {value.id for value in DSFValue.query.filter_by(dsf_id=dsf.id)}
    assert rendered_value_ids == stored_value_ids
    assert len(rendered_value_ids) == FULL_DSF_COLUMN_COUNT
    assert compact_question_rows == 31

    target_values = (
        DSFValue.query.join(ImportColumn)
        .filter(
            DSFValue.dsf_id == dsf.id,
            ImportColumn.variable_name.ilike("Aménagement, agencements et installations%"),
        )
        .order_by(ImportColumn.column_index)
        .all()
    )
    assert [value.column.column_index for value in target_values] == [61, 62, 63, 64]
    assert [value.column.variable_name for value in target_values] == headers[60:64]

    for index, value in enumerate(target_values, start=1):
        update = client.patch(
            f"/dsf/api/values/{value.id}",
            json={"value": str(index * 100), "status": "verified"},
        )
        assert update.status_code == 200
        assert update.get_json()["raw_value"] == str(index * 100)
        db.session.refresh(value)
        assert raw_input_value(value.current_value) == str(index * 100)
