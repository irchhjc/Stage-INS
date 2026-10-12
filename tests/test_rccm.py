import pytest
from openpyxl import load_workbook
from sqlalchemy import inspect, text

from app.extensions import db
from app.models import AuditLog, DSF
from app.services.auth_service import create_user
from app.services.database_service import upgrade_legacy_schema
from app.services.dsf_service import clean_rccm, search_dsfs
from app.services.export_service import export_all_completed_workbooks, export_controlled_workbook


def _first_dsf():
    return DSF.query.order_by(DSF.id).first()


def test_clean_rccm_accepts_anything_without_length_or_format_limit():
    assert clean_rccm("  cg-bzv  01/2020  b12 ") == "cg-bzv  01/2020  b12"  # saisie conservée telle quelle
    assert clean_rccm("RC<script>é€ #1 — N°7") == "RC<script>é€ #1 — N°7"
    long_value = "A" * 5000
    assert clean_rccm(long_value) == long_value
    assert clean_rccm("") is None and clean_rccm("   ") is None and clean_rccm(None) is None
    assert clean_rccm("AB\x00CD") == "ABCD"  # seul le caractère NUL (non stockable) est retiré


def test_admin_saves_rccm_and_the_change_is_audited(client, imported_session):
    dsf = _first_dsf()

    response = client.patch(f"/dsf/api/{dsf.id}/rccm", json={"value": "CG-BZV-01-2020-B12-00345"})

    assert response.status_code == 200
    assert response.get_json() == {"ok": True, "rccm": "CG-BZV-01-2020-B12-00345"}
    assert db.session.get(DSF, dsf.id).rccm == "CG-BZV-01-2020-B12-00345"
    log = AuditLog.query.filter_by(dsf_id=dsf.id, action="modification RCCM").one()
    assert (log.old_value, log.new_value, log.operator) == (None, "CG-BZV-01-2020-B12-00345", "irch")

    # Une valeur identique n'ajoute pas d'entrée de journal ; vide = effacer.
    client.patch(f"/dsf/api/{dsf.id}/rccm", json={"value": "CG-BZV-01-2020-B12-00345"})
    assert AuditLog.query.filter_by(dsf_id=dsf.id, action="modification RCCM").count() == 1
    assert client.patch(f"/dsf/api/{dsf.id}/rccm", json={"value": ""}).get_json()["rccm"] == ""
    assert db.session.get(DSF, dsf.id).rccm is None


def test_any_text_of_any_length_is_saved_as_typed(client, imported_session):
    dsf = _first_dsf()
    free_text = "RC/12 <b>é</b> — n°7 " + "x" * 3000

    response = client.patch(f"/dsf/api/{dsf.id}/rccm", json={"value": free_text})

    assert response.status_code == 200
    assert db.session.get(DSF, dsf.id).rccm == free_text
    page = client.get(f"/dsf/{dsf.id}/fiche/IDENT").get_data(as_text=True)
    assert "&lt;b&gt;" in page and "<b>é</b>" not in page  # affiché échappé, jamais interprété
    assert 'maxlength' not in page.split('id="rccmInput"')[1].split(">")[0]


def test_only_the_assigned_controller_can_edit_the_rccm(app, client, imported_session):
    dsf = _first_dsf()
    owner = create_user("proprietaire", "motdepasse10", role="controller")
    create_user("intrus", "motdepasse10", role="controller")
    dsf.assigned_to_id = owner.id
    db.session.commit()

    def as_user(name):
        other = app.test_client()
        other.post("/auth/login", data={"username": name, "password": "motdepasse10"})
        return other.patch(f"/dsf/api/{dsf.id}/rccm", json={"value": "RC 1"})

    assert as_user("intrus").status_code == 403
    assert as_user("proprietaire").status_code == 200


def test_identification_page_shows_the_rccm_field_but_other_fiches_do_not(client, imported_session):
    dsf = _first_dsf()
    dsf.rccm = "CG-BZV-01-2020-B12-00345"
    db.session.commit()

    ident = client.get(f"/dsf/{dsf.id}/fiche/IDENT").get_data(as_text=True)
    other = client.get(f"/dsf/{dsf.id}/fiche/BILAN_ACTIF").get_data(as_text=True)

    assert 'id="rccmInput"' in ident and 'value="CG-BZV-01-2020-B12-00345"' in ident
    assert "Registre du commerce (RCCM)" in ident
    assert "RCCM CG-BZV-01-2020-B12-00345" in ident
    assert 'id="rccmInput"' not in other
    assert "RCCM CG-BZV-01-2020-B12-00345" in client.get(f"/dsf/{dsf.id}").get_data(as_text=True)


def test_dsfs_are_searchable_by_rccm_even_without_separators(client, imported_session):
    first, second = DSF.query.order_by(DSF.id).all()
    first.rccm = "CG-BZV-01-2020-B12-00345"
    db.session.commit()

    assert [d.id for d in search_dsfs("cgbzv012020")] == [first.id]
    assert [d.id for d in search_dsfs("b12-00345")] == [first.id]
    assert second.id not in [d.id for d in search_dsfs("CG-BZV")]


def test_existing_database_gets_the_rccm_column_on_startup(app):
    with db.engine.begin() as connection:
        connection.execute(text("ALTER TABLE dsfs DROP COLUMN rccm"))
    assert "rccm" not in {c["name"] for c in inspect(db.engine).get_columns("dsfs")}

    upgrade_legacy_schema()

    assert "rccm" in {c["name"] for c in inspect(db.engine).get_columns("dsfs")}


def test_rccm_is_included_in_exports(client, imported_session):
    first, second = DSF.query.order_by(DSF.row_index).all()
    first.status = "completed"
    first.rccm = "CG-BZV-01-2020-B12-00345"
    db.session.commit()

    path, _, _ = export_all_completed_workbooks()
    book = load_workbook(path)
    sheet = book["DSF_TERMINEES"]
    assert sheet.cell(1, sheet.max_column).value == "RCCM (saisi dans l'application)"
    assert sheet.cell(2, sheet.max_column).value == "CG-BZV-01-2020-B12-00345"
    book.close()

    path = export_controlled_workbook(imported_session.id)
    book = load_workbook(path)
    sheet = book[imported_session.sheet_name]
    assert sheet.cell(imported_session.header_row, sheet.max_column).value == "RCCM (saisi dans l'application)"
    assert sheet.cell(imported_session.header_row + 1, sheet.max_column).value == "CG-BZV-01-2020-B12-00345"
    book.close()


def test_very_long_rccm_does_not_break_the_exports(client, imported_session):
    first = DSF.query.order_by(DSF.row_index).first()
    first.status = "completed"
    first.rccm = "Z" * 40000
    db.session.commit()

    path, _, _ = export_all_completed_workbooks()
    book = load_workbook(path)
    sheet = book["DSF_TERMINEES"]
    assert len(sheet.cell(2, sheet.max_column).value) == 32767
    book.close()
