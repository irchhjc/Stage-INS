"""Génère app/config/fiche_layouts.py à partir du classeur modèle.

Usage : python scripts/build_fiche_layouts.py [classeur.xlsx]

Le classeur (par défaut fiche_en_tableau/ouioui.xlsx) décrit comment chaque
fiche doit être présentée : un bandeau marine par fiche, une ligne d'en-têtes
bleu clair (« Poste », puis les colonnes) et une ligne par poste. Seule la
structure est extraite (libellés et nombre de colonnes) ; les lettres de
colonnes du classeur ne sont pas utilisées, car les valeurs sont rattachées
par position au sein de chaque fiche détectée à l'import.
"""
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_SOURCE = ROOT / "fiche_en_tableau" / "ouioui.xlsx"
TARGET = ROOT / "app" / "config" / "fiche_layouts.py"

TITLE_FILL = "FF1F3864"
HEADER_FILL = "FFD9E1F2"
QUESTION_PREFIX = "Note renseign"

# Blocs du classeur -> code de fiche de l'application (dans l'ordre du classeur).
BLOCK_CODES = [
    ("IDENTIFICATION", "IDENT"),
    ("BILAN – ACTIF", "BILAN_ACTIF"),
    ("BILAN – PASSIF", "BILAN_PASSIF"),
    ("COMPTE DE RÉSULTAT", "COMPTE_RESULTAT"),
    ("NOTE 3A", "NOTE_3A"),
    ("NOTE 3C", "NOTE_3A"),
    ("NOTE 3D", "NOTE_AMORT"),
    ("NOTE 4", "NOTE_4"),
    ("NOTE 5 (ACTIF", "NOTE_5_ACTIF"),
    ("NOTE 5 (DETTES", "NOTE_5_PASSIF"),
    ("NOTE 6", "NOTE_6"),
    ("NOTE 7", "NOTE_7"),
    ("NOTE 8", "NOTE_8"),
    ("NOTE 9", "NOTE_9"),
    ("NOTE 10", "NOTE_10"),
    ("NOTE 11", "NOTE_11"),
    ("NOTE 14", "NOTE_14"),
    ("NOTE 15A", "NOTE_15A"),
    ("NOTE 15B", "NOTE_15B"),
    ("NOTE 16A", "NOTE_16A"),
    ("NOTE 16B", "NOTE_16B"),
    ("NOTE 17", "NOTE_17"),
    ("NOTE 18", "NOTE_18"),
    ("NOTE 19", "NOTE_19"),
    ("NOTE 20", "NOTE_20"),
    ("NOTE 21", "NOTE_21"),
    ("NOTE 22", "NOTE_22"),
    ("NOTE 23", "NOTE_23"),
    ("NOTE 24", "NOTE_24"),
    ("NOTE 25", "NOTE_25"),
    ("NOTE 26", "NOTE_26"),
    ("NOTE 27A", "NOTE_27A"),
    ("NOTE 27B", "NOTE_27B"),
    ("NOTE 29", "NOTE_29"),
    ("NOTE 34", "NOTE_34"),
    ("INFORMATIONS", "TRACE"),
]


def _is_optional(code, block_title, label):
    """Lignes présentes dans certaines versions seulement du fichier source."""
    if code == "IDENT" and label == "Cle":
        return True
    return (
        code == "NOTE_3A"
        and block_title.upper().startswith("NOTE 3C")
        and label.startswith(QUESTION_PREFIX)
    )


def _fill(cell):
    return cell.fill.fgColor.rgb if cell.fill.fill_type else None


def _code_for(title):
    upper = str(title).upper()
    for prefix, code in BLOCK_CODES:
        if upper.startswith(prefix):
            return code
    raise ValueError(f"Bloc inconnu dans le classeur : {title!r}")


def extract(path):
    sheet = openpyxl.load_workbook(path).active
    layouts, block_counts = {}, {}
    block_title = code = table = None
    pending_title = None

    def new_table(columns, question=False):
        nonlocal table, pending_title
        table = {"title": pending_title, "columns": columns, "question": question, "rows": []}
        pending_title = None
        layouts.setdefault(code, []).append(table)

    for number in range(7, sheet.max_row + 1):
        label = sheet.cell(number, 2).value
        if label is None:
            continue
        label = str(label).strip()
        fill = _fill(sheet.cell(number, 2))
        if fill == TITLE_FILL:
            block_title, code = label, _code_for(label)
            block_counts[code] = block_counts.get(code, 0) + 1
            # Une fiche qui regroupe plusieurs blocs (3A + 3C) intitule les suivants.
            pending_title = label.title() if block_counts[code] > 1 else None
            table = None
            inline_header = str(sheet.cell(number, 3).value or "").strip() == "valeur"
            if inline_header or code == "TRACE":
                new_table(["Valeur"])
            continue
        if fill == HEADER_FILL:
            new_table(
                [
                    str(sheet.cell(number, column).value).strip()
                    for column in range(3, 10)
                    if sheet.cell(number, column).value is not None
                ]
            )
            continue
        question = label.startswith(QUESTION_PREFIX)
        if question:
            new_table(["Valeur"], question=True)
        row = {"label": label}
        if _is_optional(code, block_title, label):
            row["optional"] = True
        table["rows"].append(row)
        if question:
            table = None
    return layouts


def render(layouts):
    lines = [
        '"""Structure des fiches DSF en tableaux.',
        "",
        "Fichier généré par scripts/build_fiche_layouts.py à partir de",
        "fiche_en_tableau/ouioui.xlsx. Ne pas modifier à la main : modifier le",
        "classeur puis relancer le script.",
        "",
        "Chaque fiche est une suite de tableaux ; chaque tableau a des colonnes et",
        "des lignes (postes). Les valeurs sont rattachées dans l'ordre : ligne par",
        "ligne, colonne par colonne. Une ligne « optional » n'existe que dans",
        'certaines versions du fichier source.',
        '"""',
        "",
        "FICHE_LAYOUTS = {",
    ]
    for code, tables in layouts.items():
        lines.append(f"    {code!r}: [")
        for table in tables:
            lines.append("        {")
            lines.append(f"            \"title\": {table['title']!r},")
            lines.append(f"            \"columns\": {table['columns']!r},")
            if table["question"]:
                lines.append("            \"question\": True,")
            lines.append("            \"rows\": [")
            for row in table["rows"]:
                lines.append(f"                {row!r},")
            lines.append("            ],")
            lines.append("        },")
        lines.append("    ],")
    lines.append("}")
    return "\n".join(lines) + "\n"


if __name__ == "__main__":
    source = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_SOURCE
    layouts = extract(source)
    TARGET.write_text(render(layouts), encoding="utf-8", newline="\n")
    total = sum(
        len(table["rows"]) * len(table["columns"])
        for tables in layouts.values()
        for table in tables
    )
    print(f"{len(layouts)} fiches, {total} cellules -> {TARGET}")
