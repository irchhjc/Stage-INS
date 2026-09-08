import uuid
from copy import copy
from pathlib import Path

from flask import current_app
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy.orm import joinedload

from app.extensions import db
from app.models import DSF, DSFValue, ImportColumn, ImportSession
from app.services.value_codec import deserialize_value, display_value


COLORS = {
    "unverified": "FF000000",
    "verified": "FF008000",
    "anomaly": "FFC00000",
    "not_applicable": "FF666666",
}
CORRECTED_FILL = PatternFill(fill_type="solid", fgColor="FFE2F0D9")


def _apply_font_color(cell, color):
    font = copy(cell.font) if cell.font else Font()
    font.color = color
    cell.font = font


def _journal_sheet_name(workbook):
    base = "JOURNAL_CONTROLE"
    if base not in workbook.sheetnames:
        return base
    number = 2
    while f"{base}_{number}" in workbook.sheetnames:
        number += 1
    return f"{base}_{number}"


def export_controlled_workbook(import_session_id, assigned_user_id=None, username=None):
    import_session = db.session.get(ImportSession, import_session_id)
    if import_session is None:
        raise ValueError("Session d'import introuvable.")

    completed_dsfs_query = DSF.query.filter(
        DSF.import_session_id == import_session.id,
        DSF.status == "completed",
    )
    if assigned_user_id is not None:
        completed_dsfs_query = completed_dsfs_query.filter(DSF.assigned_to_id == assigned_user_id)
    completed_dsfs = completed_dsfs_query.order_by(DSF.row_index).all()
    if not completed_dsfs:
        raise ValueError("Aucune DSF entièrement contrôlée n'est disponible pour l'export.")

    completed_dsf_ids = [dsf.id for dsf in completed_dsfs]
    selected_rows = {dsf.row_index for dsf in completed_dsfs}
    source = Path(import_session.filepath)
    if not source.exists():
        raise FileNotFoundError("Le classeur original associé à cet import est introuvable.")

    workbook = load_workbook(source, data_only=False, keep_links=True)
    if import_session.sheet_name not in workbook.sheetnames:
        raise ValueError("La feuille source enregistrée n'existe plus dans le classeur original.")
    worksheet = workbook[import_session.sheet_name]

    values_query = (
        DSFValue.query.options(joinedload(DSFValue.dsf), joinedload(DSFValue.column))
        .join(DSF)
        .join(ImportColumn)
        .filter(
            DSF.import_session_id == import_session.id,
            DSF.id.in_(completed_dsf_ids),
        )
    )
    values = values_query.order_by(DSF.row_index, ImportColumn.column_index).all()
    for value in values:
        cell = worksheet.cell(row=value.dsf.row_index, column=value.column.column_index)
        if value.corrected:
            cell.value = deserialize_value(value.current_value)
        _apply_font_color(cell, COLORS.get(value.status, COLORS["unverified"]))
        if value.corrected and value.status == "verified":
            cell.fill = copy(CORRECTED_FILL)

    for row_index in range(worksheet.max_row, import_session.header_row, -1):
        if row_index not in selected_rows:
            worksheet.delete_rows(row_index)

    if assigned_user_id is not None:
        for other_sheet in list(workbook.worksheets):
            if other_sheet is not worksheet:
                workbook.remove(other_sheet)

    journal = workbook.create_sheet(_journal_sheet_name(workbook))
    headers = [
        "NIU",
        "NUMERO DSF",
        "Raison sociale",
        "Année",
        "Fiche",
        "Variable",
        "Colonne Excel",
        "Valeur originale",
        "Nouvelle valeur",
        "Statut",
        "Opérateur",
        "Date",
        "Heure",
    ]
    journal.append(headers)
    header_fill = PatternFill(fill_type="solid", fgColor="FF17365D")
    for cell in journal[1]:
        cell.font = Font(color="FFFFFFFF", bold=True)
        cell.fill = header_fill

    for value in values:
        if value.status == "unverified" and not value.corrected:
            continue
        timestamp = value.corrected_at or value.verified_at
        journal.append(
            [
                value.dsf.niu,
                value.dsf.numero_dsf,
                value.dsf.raison_sociale,
                value.dsf.annee,
                value.column.fiche_name,
                value.variable_name,
                value.column.excel_column_letter,
                display_value(value.original_value),
                display_value(value.current_value),
                "corrigée et vérifiée" if value.corrected and value.verified else value.status,
                value.operator,
                timestamp.date().isoformat() if timestamp else None,
                timestamp.time().replace(microsecond=0).isoformat() if timestamp else None,
            ]
        )
    journal.freeze_panes = "A2"
    journal.auto_filter.ref = journal.dimensions
    widths = [20, 18, 32, 12, 38, 55, 15, 22, 22, 24, 22, 14, 12]
    for index, width in enumerate(widths, start=1):
        journal.column_dimensions[journal.cell(1, index).column_letter].width = width

    output_dir = Path(current_app.config["EXPORT_FOLDER"])
    output_dir.mkdir(parents=True, exist_ok=True)
    safe_stem = Path(import_session.filename).stem[:80] or "dsf"
    scope = f"_{username}" if username else ""
    output_path = output_dir / f"{safe_stem}_dsf_controlees{scope}_{uuid.uuid4().hex[:8]}.xlsx"
    workbook.save(output_path)
    workbook.close()
    return output_path
