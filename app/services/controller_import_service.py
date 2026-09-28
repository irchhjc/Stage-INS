import io
import unicodedata
from dataclasses import dataclass
from zipfile import BadZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils.exceptions import InvalidFileException

from app.extensions import db
from app.models import User
from app.services.auth_service import normalize_full_name, validate_new_user


MAX_CONTROLLER_ROWS = 1_000
TEMPLATE_HEADERS = ("NOM COMPLET", "NOM UTILISATEUR", "MOT DE PASSE")


class ControllerImportError(ValueError):
    pass


@dataclass(frozen=True)
class ControllerImportResult:
    created_count: int
    rejected_errors: tuple[str, ...]
    ignored_blank_rows: int

    @property
    def rejected_count(self):
        return len(self.rejected_errors)


def _normalize_header(value):
    text = unicodedata.normalize("NFKD", str(value or ""))
    text = "".join(character for character in text if not unicodedata.combining(character))
    return " ".join(
        text.casefold()
        .replace("’", "'")
        .replace("_", " ")
        .replace("-", " ")
        .split()
    )


HEADER_ALIASES = {
    "full_name": {
        _normalize_header(value)
        for value in ("NOM COMPLET", "NOM ET PRENOM", "NOM ET PRÉNOM", "FULL NAME")
    },
    "username": {
        _normalize_header(value)
        for value in ("NOM UTILISATEUR", "NOM D'UTILISATEUR", "IDENTIFIANT", "USERNAME")
    },
    "password": {
        _normalize_header(value)
        for value in ("MOT DE PASSE", "PASSWORD")
    },
}


def _column_mapping(header_values):
    mapping = {}
    normalized_headers = [_normalize_header(value) for value in header_values]
    for field, aliases in HEADER_ALIASES.items():
        matches = [index for index, header in enumerate(normalized_headers) if header in aliases]
        if len(matches) > 1:
            raise ControllerImportError(
                f"La colonne « {TEMPLATE_HEADERS[list(HEADER_ALIASES).index(field)]} » apparaît plusieurs fois."
            )
        if matches:
            mapping[field] = matches[0]

    missing = [
        TEMPLATE_HEADERS[list(HEADER_ALIASES).index(field)]
        for field in HEADER_ALIASES
        if field not in mapping
    ]
    if missing:
        raise ControllerImportError("Colonne(s) obligatoire(s) absente(s) : " + ", ".join(missing) + ".")
    return mapping


def _cell(row, index):
    return row[index] if index < len(row) else None


def import_controllers(source):
    try:
        workbook = load_workbook(source, read_only=True, data_only=True)
    except (InvalidFileException, BadZipFile, OSError, ValueError) as exc:
        raise ControllerImportError("Le fichier Excel est illisible ou invalide.") from exc

    try:
        worksheet = workbook.active
        rows = worksheet.iter_rows(values_only=True)
        header = next(rows, None)
        if header is None:
            raise ControllerImportError("Le fichier Excel est vide.")
        mapping = _column_mapping(header)

        existing_usernames = {username for (username,) in db.session.query(User.username).all()}
        seen_usernames = set()
        valid_users = []
        errors = []
        ignored_blank_rows = 0

        for excel_row_number, row in enumerate(rows, start=2):
            if excel_row_number > MAX_CONTROLLER_ROWS + 1:
                raise ControllerImportError(
                    f"Le fichier dépasse la limite de {MAX_CONTROLLER_ROWS} contrôleurs par import."
                )

            full_name_value = _cell(row, mapping["full_name"])
            username_value = _cell(row, mapping["username"])
            password_value = _cell(row, mapping["password"])
            if all(value is None or str(value).strip() == "" for value in (full_name_value, username_value, password_value)):
                ignored_blank_rows += 1
                continue

            try:
                if not isinstance(full_name_value, str):
                    raise ValueError("le nom complet doit être du texte")
                full_name = normalize_full_name(full_name_value)
                if not full_name:
                    raise ValueError("le nom complet est obligatoire")
                if not isinstance(username_value, str):
                    raise ValueError("le nom d'utilisateur doit être du texte")
                if not isinstance(password_value, str):
                    raise ValueError("le mot de passe doit être au format texte")
                username = validate_new_user(username_value, password_value)
                if username in existing_usernames:
                    raise ValueError("ce nom d'utilisateur existe déjà")
                if username in seen_usernames:
                    raise ValueError("ce nom d'utilisateur est répété dans le fichier")
            except ValueError as exc:
                errors.append(f"Ligne {excel_row_number} : {exc}.")
                continue

            user = User(username=username, full_name=full_name, role="controller")
            user.set_password(password_value)
            valid_users.append(user)
            seen_usernames.add(username)

        if not valid_users and not errors:
            raise ControllerImportError("Le fichier ne contient aucun contrôleur à créer.")

        try:
            db.session.add_all(valid_users)
            db.session.commit()
        except Exception:
            db.session.rollback()
            raise

        return ControllerImportResult(
            created_count=len(valid_users),
            rejected_errors=tuple(errors),
            ignored_blank_rows=ignored_blank_rows,
        )
    finally:
        workbook.close()


def build_controller_import_template():
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "CONTROLEURS"
    worksheet.append(TEMPLATE_HEADERS)
    worksheet.freeze_panes = "A2"
    fill = PatternFill("solid", fgColor="1F4E78")
    for cell in worksheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = fill
    worksheet.column_dimensions["A"].width = 32
    worksheet.column_dimensions["B"].width = 24
    worksheet.column_dimensions["C"].width = 24

    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    output.seek(0)
    return output
