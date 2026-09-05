import shutil
import uuid
from pathlib import Path

from flask import current_app
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from sqlalchemy import insert

from app.config.fiche_mapping import FICHE_DEFINITIONS
from app.extensions import db
from app.models import DSF, DSFValue, FicheStatus, ImportColumn, ImportSession
from app.services.mapping_service import MappingError, build_column_mapping
from app.services.value_codec import deserialize_value, display_value, serialize_value


IDENTITY_HEADERS = {
    "numero_dsf": "NUMERO DE LA DSF",
    "numero_dsf_t": "NUMERO DSF_T",
    "niu": "NIU",
    "cle": "Cle",
    "raison_sociale": "Raison sociale",
    "sigle": "Sigle usuel",
    "annee": "Année",
}


class ExcelImportError(ValueError):
    pass


def _sheet_candidates(workbook):
    candidates = []
    for worksheet in workbook.worksheets:
        for row_index, values in enumerate(
            worksheet.iter_rows(min_row=1, max_row=min(worksheet.max_row, 20), values_only=True),
            start=1,
        ):
            if "NIU" in values:
                candidates.append((worksheet.title, row_index))
    return candidates


def inspect_workbook(filepath, sheet_name=None, header_row=None):
    try:
        workbook = load_workbook(filepath, read_only=True, data_only=False)
    except Exception as exc:
        raise ExcelImportError(f"Le fichier n'est pas un classeur .xlsx lisible : {exc}") from exc

    try:
        if sheet_name:
            if sheet_name not in workbook.sheetnames:
                raise ExcelImportError(f"La feuille « {sheet_name} » n'existe pas.")
            worksheet = workbook[sheet_name]
            selected_header_row = int(header_row or 1)
        else:
            candidates = _sheet_candidates(workbook)
            unique_candidates = list(dict.fromkeys(candidates))
            if len(unique_candidates) != 1:
                details = ", ".join(f"{name} (ligne {row})" for name, row in unique_candidates) or "aucune"
                raise ExcelImportError(
                    "Impossible de choisir silencieusement la feuille et la ligne d'en-tête. "
                    f"Candidatures contenant NIU : {details}."
                )
            sheet_name, selected_header_row = unique_candidates[0]
            worksheet = workbook[sheet_name]

        headers = list(
            next(
                worksheet.iter_rows(
                    min_row=selected_header_row,
                    max_row=selected_header_row,
                    max_col=worksheet.max_column,
                    values_only=True,
                )
            )
        )
        if not headers or all(value is None for value in headers):
            raise ExcelImportError("La ligne d'en-tête est vide.")
        if headers.count("NIU") != 1:
            raise ExcelImportError("La colonne exacte « NIU » doit apparaître une seule fois.")
        try:
            mapping = build_column_mapping(headers)
        except MappingError as exc:
            raise ExcelImportError(str(exc)) from exc

        return {
            "sheet_name": worksheet.title,
            "header_row": selected_header_row,
            "max_row": worksheet.max_row,
            "max_column": worksheet.max_column,
            "headers": headers,
            "mapping": mapping,
        }
    finally:
        workbook.close()


def save_uploaded_file(file_storage):
    upload_root = Path(current_app.config["UPLOAD_FOLDER"])
    session_folder = upload_root / uuid.uuid4().hex
    session_folder.mkdir(parents=True, exist_ok=False)
    destination = session_folder / "original.xlsx"
    file_storage.save(destination)
    return destination


def copy_local_workbook(source_path):
    upload_root = Path(current_app.config["UPLOAD_FOLDER"])
    session_folder = upload_root / uuid.uuid4().hex
    session_folder.mkdir(parents=True, exist_ok=False)
    destination = session_folder / "original.xlsx"
    shutil.copy2(source_path, destination)
    return destination


def import_workbook(filepath, original_filename, sheet_name=None, header_row=None):
    metadata = inspect_workbook(filepath, sheet_name, header_row)
    workbook = load_workbook(filepath, read_only=True, data_only=False)
    worksheet = workbook[metadata["sheet_name"]]

    import_session = ImportSession(
        filename=original_filename,
        filepath=str(Path(filepath).resolve()),
        sheet_name=metadata["sheet_name"],
        header_row=metadata["header_row"],
        row_count=0,
        column_count=metadata["max_column"],
    )
    db.session.add(import_session)
    db.session.flush()

    db.session.execute(
        insert(ImportColumn),
        [
            {
                "import_session_id": import_session.id,
                "column_index": mapped["column_index"],
                "excel_column_letter": get_column_letter(mapped["column_index"]),
                "variable_name": mapped["variable_name"],
                "fiche_code": mapped["fiche_code"],
                "fiche_name": mapped["fiche_name"],
            }
            for mapped in metadata["mapping"]
        ],
    )
    columns = (
        ImportColumn.query.filter_by(import_session_id=import_session.id)
        .order_by(ImportColumn.column_index)
        .all()
    )

    header_positions = {}
    for field, header in IDENTITY_HEADERS.items():
        matches = [index + 1 for index, value in enumerate(metadata["headers"]) if value == header]
        header_positions[field] = matches[0] if len(matches) == 1 else None

    imported_rows = 0
    try:
        rows = worksheet.iter_rows(
            min_row=metadata["header_row"] + 1,
            max_row=metadata["max_row"],
            max_col=metadata["max_column"],
            values_only=True,
        )
        for row_index, row_tuple in enumerate(rows, start=metadata["header_row"] + 1):
            row_values = list(row_tuple)
            if all(value is None or (isinstance(value, str) and not value.strip()) for value in row_values):
                continue

            identity = {}
            for field, column_index in header_positions.items():
                raw = row_values[column_index - 1] if column_index else None
                identity[field] = None if raw is None else str(raw).strip()

            dsf = DSF(import_session_id=import_session.id, row_index=row_index, **identity)
            if not dsf.niu:
                dsf.anomaly_count = 1
            db.session.add(dsf)
            db.session.flush()

            dsf_values = []
            for column, raw_value in zip(columns, row_values):
                serialized = serialize_value(raw_value)
                dsf_values.append(
                    {
                        "dsf_id": dsf.id,
                        "import_column_id": column.id,
                        "variable_name": column.variable_name,
                        "original_value": serialized,
                        "current_value": serialized,
                        "status": "unverified",
                        "verified": False,
                        "corrected": False,
                    }
                )
            db.session.execute(insert(DSFValue), dsf_values)
            db.session.execute(
                insert(FicheStatus),
                [
                    {
                        "dsf_id": dsf.id,
                        "fiche_code": definition["code"],
                        "fiche_name": definition["name"],
                        "position": position,
                        "status": "not_started",
                    }
                    for position, definition in enumerate(FICHE_DEFINITIONS, start=1)
                ],
            )
            imported_rows += 1

        import_session.row_count = imported_rows
        db.session.commit()
        from app.services.dsf_service import recompute_dsf_progress

        for dsf in DSF.query.filter_by(import_session_id=import_session.id).all():
            recompute_dsf_progress(dsf)
        db.session.commit()
    except Exception:
        db.session.rollback()
        raise
    finally:
        workbook.close()

    return import_session


def preview_identity(dsf):
    return {
        "NIU": dsf.niu or "—",
        "NUMERO DE LA DSF": dsf.numero_dsf or "—",
        "Raison sociale": dsf.raison_sociale or "—",
        "Sigle": dsf.sigle or "—",
        "Année": dsf.annee or "—",
    }


def display_current_value(dsf_value):
    return display_value(dsf_value.current_value)


def excel_current_value(dsf_value):
    return deserialize_value(dsf_value.current_value)
